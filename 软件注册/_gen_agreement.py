# -*- coding: utf-8 -*-
"""生成《合作开发协议》docx（软件开发方为多人时使用）。

注意：本协议仅在"开发方式 = 合作开发"时需要提交。若软件为一人独立完成，
申请表里的开发方式应填"独立开发"，本文件不必提交。
"""
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

OUTDIR = pathlib.Path(r'F:\MolecuLens\软件注册')
SOFT_FULL = 'MolecuLens 分子结构与轨迹显示引擎'
SOFT_SHORT = 'MolecuLens'
VERSION = 'V1.0'

# 参与开发的各方：填真实姓名与身份证号后再生成
PARTIES = [
    ('甲方', '【　　】', '【　　　　　　　　　　】'),
    ('乙方', '【　　】', '【　　　　　　　　　　】'),
]


def P(text, indent=True, align=None):
    p = doc.add_paragraph(text)
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    if align is not None:
        p.alignment = align
    return p


doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Emu(7560310), Emu(10692130)
s.top_margin = s.bottom_margin = Emu(914400)
s.left_margin = s.right_margin = Emu(1143000)

st = doc.styles['Normal']
st.font.name = '宋体'
st.font.size = Pt(12)
st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
st.paragraph_format.line_spacing = 1.6
st.paragraph_format.space_after = Pt(4)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('合 作 开 发 协 议')
r.font.size = Pt(20)
r.font.name = '黑体'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
title.paragraph_format.space_after = Pt(18)

P(f'协议各方就合作开发“{SOFT_FULL}”（以下简称“本软件”）事宜，'
  '经友好协商，达成如下协议：')

doc.add_heading('一、开发内容', level=2)
P(f'各方共同开发本软件 {VERSION}，软件功能与版本以《软件著作权申请表》'
  '所载内容为准。开发完成日期为 2026 年 9 月 15 日。')

doc.add_heading('二、各方分工', level=2)
for role, name, _ in PARTIES:
    P(f'{role}（{name}）：承担本软件相应模块的设计、编码与调试工作，'
      '并配合完成整体联调与测试。', indent=True)

doc.add_heading('三、著作权归属', level=2)
P('本软件的著作权由各方共同享有，各方均为著作权人。'
  '任何一方对外转让或许可他人使用本软件，须经各方书面同意。')

doc.add_heading('四、权利行使', level=2)
P('各方均可为自身教学、科研或非商业目的自由使用本软件，无需另行征得同意。'
  '涉及商业用途的，由各方另行协商确定收益分配办法。')

doc.add_heading('五、保密条款', level=2)
P('各方对在合作开发过程中知悉的对方技术资料与未公开信息负有保密义务，'
  '未经对方书面同意不得向第三方披露。')

doc.add_heading('六、其他', level=2)
P('本协议未尽事宜，由各方协商解决。本协议一式【　】份，各方各执一份，'
  '具有同等效力，自各方签字之日起生效。')

doc.add_paragraph()
doc.add_paragraph()

# 落款
for role, name, idno in PARTIES:
    p = doc.add_paragraph(f'{role}（签字）：{name}')
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(f'身份证号：{idno}')
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph('日　　期：　　年　　月　　日')
    p.paragraph_format.space_after = Pt(16)
    doc.add_paragraph()

out = OUTDIR / f'{SOFT_SHORT}_合作开发协议.docx'
doc.save(out)
print('已生成', out)
print('  待填：各方姓名、身份证号、协议份数、签字日期')
