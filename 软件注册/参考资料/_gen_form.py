# -*- coding: utf-8 -*-
"""生成《软件著作权申请表》xlsx。"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

OUTDIR = Path(r'H:\2026数学建模\cpp\软著注册')
SOFT_FULL = '药材烘干过程数值模拟软件'
SOFT_SHORT = '烘干模拟'
VERSION = 'V1.0'

ROWS = [
    ('软件基本信息', '软件全称', SOFT_FULL),
    ('', '软件简称', SOFT_SHORT),
    ('', '版本号', VERSION),
    ('', '软件分类', '应用软件'),
    ('', '软件说明', '原创'),
    ('开发方式', '', '合作开发'),
    ('开发完成日期', '', '2026-09-15'),
    ('发表状态', '', '未发表'),
    ('软件功能与特点', '开发的硬件环境',
     '内存：16GB  硬盘：1TB'),
    ('', '运行的硬件环境', '内存：2GB  硬盘：50MB'),
    ('', '开发该软件的操作系统', 'Windows 11'),
    ('', '软件开发环境 / 开发工具',
     'Visual Studio Code、PyCharm、MinGW-w64 g++、CMake'),
    ('', '该软件的运行平台 / 操作系统', 'Windows 10 / 11（64 位）'),
    ('', '软件运行支撑环境 / 支持软件',
     '无。核心求解器静态链接，运行时不依赖第三方运行库；'
     '图形界面需 Python 3.9 及以上，并安装 numpy、scipy、matplotlib、'
     'openpyxl、scikit-learn'),
    ('', '编程语言', 'C++、Python'),
    ('', '源程序量', 7687),
    ('', '开发目的',
     '针对药材烘干过程中水分在物料内部迁移缓慢、烘干时长难以准确预估的问题，'
     '提供基于数值求解的烘干过程模拟与时长预测工具，辅助确定烘干工艺参数，'
     '减少因预估偏短造成的中心部位未达标。'),
    ('', '面向领域 / 行业',
     '面向中药材加工与农产品干燥行业，服务烘干工艺技术人员，'
     '适配无限长圆柱类多孔物料的传热传质计算；也可作为相关专业教学辅助工具，'
     '支持通过命令行批量调用与二次开发。'),
    ('', '软件的主要功能',
     '本软件以烘房温湿度曲线和物料半径数据为输入，对物料内部的温度场与'
     '含水率场进行一维径向数值模拟。软件按用途分为四个功能模块：'
     '预热阶段模拟，给出升温与表面失水的短时段过程；全周期干燥模拟，'
     '给出完整烘干过程的场演化；干燥时长预测，推送到物料各点达标并给出'
     '所需时间；收缩工况模拟，处理半径随失水收缩的情形并分离物性与几何'
     '两种效应。软件集成网格划分、三套物性方案、两阶段液膜表面模型、'
     '湿球温度求解、终点判据定位与守恒校核等公共模块，提供图形界面与'
     '命令行两种使用方式。图形界面支持参数在线修改、计算进度实时显示、'
     '场演化云图与干燥曲线绘制；结果可按表格格式导出。'
     '内置三层精度校核，可对计算结果进行独立验证。'),
    ('', '软件的技术特点',
     '核心求解器以 C++17 编写并静态编译，单个模块的计算耗时在 20 秒以内，'
     '且不依赖任何第三方运行库，可脱离界面单独运行。'
     '数值上采用节点中心有限体积法与 Crank–Nicolson 格式，界面系数取算术平均，'
     '以适应跨两个数量级的变扩散系数；非线性用 Picard 迭代处理，'
     '三对角方程组严格对角占优。表面处理区分预热段与恒温段：预热段按相变份额'
     '折减蒸发潜热，恒温段以罚函数对流边界将表面钉在空气湿球温度，'
     '并设有液膜存在性判据，表面变干后自动退回对流边界。'
     '收缩工况采用材料坐标变换处理移动边界，网格运动项与守恒恒等式严格对应，'
     '半径收缩设为零时可精确退化为固定网格求解器。'
     '守恒量在推进循环内逐步累加，与求解器使用同一离散式，'
     '质量与能量恒等式的相对残差稳定在 10⁻¹¹ 量级。'
     '软件对保形插值、湿球温度求根等关键算法做了与通用科学计算库的逐点比对，'
     '并配有 29 项自检。'),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sheet1'

thin = Side(style='thin', color='808080')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws.merge_cells('A1:C1')
c = ws['A1']
c.value = '软件著作权申请表'
c.font = Font(name='宋体', size=14, bold=True)
c.alignment = CTR

r = 2
for a, b, v in ROWS:
    ws.cell(r, 1, a or None)
    ws.cell(r, 2, b or None)
    ws.cell(r, 3, v)
    r += 1

# 合并「分类」列中的纵向单元格
for a, b in ((2, 6), (7, 7), (8, 8), (9, 9), (10, r - 1)):
    if b > a:
        ws.merge_cells(f'A{a}:A{b}')

# ⚠ 不要把「开发方式 / 开发完成日期 / 发表状态」三行的 B:C 合并 ——
#   openpyxl 合并时只保留左上角单元格的值，合并 B:C 会把写在 C 里的内容抹掉
#   （踩过：这三行在成品里是空的）。这里本来就是"标签在 A、值在 C"。

for row in ws.iter_rows(min_row=1, max_row=r - 1, min_col=1, max_col=3):
    for cell in row:
        cell.border = BD
        if cell.row == 1:
            continue
        cell.alignment = CTR if cell.column <= 2 else WRAP
        if cell.font.size is None or cell.font.size < 10:
            cell.font = Font(name='宋体', size=10, bold=cell.column <= 2)

ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 24
ws.column_dimensions['C'].width = 86
ws.row_dimensions[1].height = 26
for i, (_, _, v) in enumerate(ROWS):
    if isinstance(v, str) and len(v) > 60:
        ws.row_dimensions[i + 2].height = 15 * (len(v) // 42 + 1)

ws.freeze_panes = 'A2'

out = OUTDIR / f'{SOFT_SHORT}_软著申请表.xlsx'
wb.save(out)
print('已生成', out)
print(f'  共 {len(ROWS)} 个字段')
