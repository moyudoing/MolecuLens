# -*- coding: utf-8 -*-
"""生成《软件著作权申请表》xlsx。"""
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

OUTDIR = pathlib.Path(r'F:\MolecuLens\软件注册')
SOFT_FULL = 'MolecuLens 分子结构与轨迹显示引擎'
SOFT_SHORT = 'MolecuLens'
VERSION = 'V1.0'

ROWS = [
    ('软件基本信息', '软件全称', SOFT_FULL),
    ('', '软件简称', SOFT_SHORT),
    ('', '版本号', VERSION),
    ('', '软件分类', '应用软件'),
    ('', '软件说明', '原创'),
    ('开发方式', '', '独立开发'),
    ('开发完成日期', '', '2026-09-15'),
    ('发表状态', '', '未发表'),
    ('软件功能与特点', '开发的硬件环境',
     '内存：16GB  硬盘：1TB  显卡：支持 Vulkan 1.1 的独立显卡'),
    ('', '运行的硬件环境',
     '内存：4GB  硬盘：200MB  显卡：支持 Vulkan 1.1'),
    ('', '开发该软件的操作系统', 'Windows 11'),
    ('', '软件开发环境 / 开发工具',
     'Visual Studio 2022、CMake 3.31、vcpkg、Vulkan SDK 1.4'),
    ('', '该软件的运行平台 / 操作系统', 'Windows 10 / 11（64 位）'),
    ('', '软件运行支撑环境 / 支持软件',
     '随程序附带 fmtd、glfw3、spdlogd 动态库，无需单独安装；'
     '运行需显卡驱动支持 Vulkan 1.1'),
    ('', '编程语言', 'C++'),
    ('', '源程序量', 27327),
    ('', '开发目的',
     '分子结构数据的原始形态是一批坐标数字，无论用于科研还是教学，'
     '第一步都要把它变成可交互的三维图形。本软件面向这一需要，'
     '提供从结构文件解析到实时三维呈现的完整链路，'
     '并把"表示方式"与"配色方案"设计为两个彼此独立的维度——'
     '任意表示可搭配任意配色，切换后立即重绘，'
     '使观察角度与信息重点的调整成为一次点击的事情。'),
    ('', '面向领域 / 行业',
     '面向结构生物学、药物化学、计算化学等科研领域，'
     '以及高等院校化学与生物类专业的教学演示；'
     '也可用于论文与讲义插图的制作。支持通过命令行批量核对结构文件，'
     '便于与既有数据处理流程衔接。'),
    ('', '软件的主要功能',
     '本软件以分子结构文件为输入，提供结构解析、三维呈现与图像输出能力。'
     '载入环节支持 PDB、GRO、SDF（V2000）三种格式，'
     '可在缺少显式键表时按原子间距自动推断键连，'
     '并识别水分子与离子等溶剂分子；解析在后台线程完成，界面不阻塞。'
     '呈现环节内置十三种表示方式（线状、棍状、球棍、CPK、VDW、甘草棍、'
     '卡通、新卡通、丝带、管状、表面、QuickSurf 等），'
     '覆盖从化学键到整体外形的各个观察尺度；'
     '并提供五种配色方案（元素 CPK、按链、按残基、按 BFactor、统一色），'
     '与表示方式自由组合。软件以"体系"为管理单位组织场景，'
     '支持显隐、重命名、整体平移、跨体系移交与独立渲染模式；'
     '具备轨道旋转、滚轮缩放、滚转与平移等完整视角操作。'
     '输出环节支持按指定分辨率与抗锯齿倍数导出 PNG 图像，'
     '并提供光栅化与光线追踪两条渲染路径。'
     '软件附带分级日志面板与命令行工具，便于排查与批量核对。'),
    ('', '软件的技术特点',
     '渲染引擎为自行实现，不依赖任何第三方渲染中间件。'
     '表示方式与配色方案在数据结构上解耦，'
     '几何生成统一输出为紧凑的命令流，再由渲染后端按批次绘制，'
     '使新增表示方式不影响渲染侧。'
     '表面与 QuickSurf 两种表示基于 Marching Cubes 等值面抽取实现，'
     '其中 QuickSurf 对原子集合构建高斯密度场并同步累加密度梯度与'
     '密度加权颜色，计算后端按 CUDA、AVX2、标量多线程三级自动选择，'
     'AVX2 路径使用八宽 SIMD 与快速指数近似，与标量结果几何一致。'
     '光照采用全局光方案：以天顶亮、地面暗的半球环境光为主体，'
     '配合均匀环境光补光，分子在任何转角下都保持可辨的明暗层次。'
     '图像导出采用超采样后盒式降采样抑制锯齿，'
     '并提供自研的 CPU 光线追踪器——按行多线程并行，'
     '对球与胶囊求解析交，光照为半球环境光并带一次球面反射，'
     '采样使用确定性抖动以保证结果可复现。'
     '配色的按链与按残基方案采用按编号哈希取色相，'
     '无需预先分配色表即可让任意链数、任意残基数下的颜色自动分散。'),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sheet1'

thin = Side(style='thin', color='808080')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws.merge_cells('A1:C1')
c = ws['A1']
c.value = '软件著作权申请表'
c.font = Font(name='宋体', size=14, bold=True)
c.alignment = CTR

r = 2
for a, b, v in ROWS:
    ws.cell(r, 1, a or None)
    ws.cell(r, 2, b or None)
    ws.cell(r, 3, v)
    r += 1

# 合并「分类」列中的纵向单元格
for a, b in ((2, 6), (7, 7), (8, 8), (9, 9), (10, r - 1)):
    if b > a:
        ws.merge_cells(f'A{a}:A{b}')

# 注意：不要把「开发方式 / 开发完成日期 / 发表状态」三行的 B:C 合并 ——
# openpyxl 合并时只保留左上角单元格的值，合并 B:C 会把写在 C 里的内容抹掉。

for row in ws.iter_rows(min_row=1, max_row=r - 1, min_col=1, max_col=3):
    for cell in row:
        cell.border = BD
        if cell.row == 1:
            continue
        cell.alignment = CTR if cell.column <= 2 else WRAP
        if cell.font.size is None or cell.font.size < 10:
            cell.font = Font(name='宋体', size=10, bold=cell.column <= 2)

ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 24
ws.column_dimensions['C'].width = 86
ws.row_dimensions[1].height = 26
for i, (_, _, v) in enumerate(ROWS):
    if isinstance(v, str) and len(v) > 60:
        ws.row_dimensions[i + 2].height = 15 * (len(v) // 42 + 1)

ws.freeze_panes = 'A2'

out = OUTDIR / f'{SOFT_SHORT}_软著申请表.xlsx'
wb.save(out)
print('已生成', out)
print(f'  共 {len(ROWS)} 个字段')
